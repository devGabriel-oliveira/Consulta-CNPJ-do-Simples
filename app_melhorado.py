"""
VIXPAR | Portal Fiscal — Consulta Simples Nacional / MEI em lote
================================================================
Consome DIRETO o endpoint /cnpj/v1 da BrasilAPI (que faz proxy para
o MinhaReceita.org) usando as chaves corretas do schema oficial:
    opcao_pelo_simples, opcao_pelo_mei,
    data_opcao_pelo_simples, data_opcao_pelo_mei,
    data_exclusao_do_simples, data_exclusao_do_mei

NÃO usa a lib mcp_fiscal_brasil aqui, porque a versão 0.5.1 procura
chaves erradas ("optante", "simples_nacional") que não existem no
retorno real da API, e por isso sempre devolve False (bug).
"""

import asyncio
import io
import re
import time
from datetime import datetime, date
from typing import Any

import httpx
import pandas as pd
import streamlit as st

# ==================== CONFIGURAÇÃO DA PÁGINA ====================
st.set_page_config(
    page_title="VIXPAR | Portal Fiscal",
    page_icon="🧾",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ==================== CUSTOM CSS ====================
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');
    html, body, [class*="css"] { font-family: 'Inter', sans-serif; }
    :root { --vix-navy: #1d2a4d; --vix-orange: #f7941d; --bg-dark: #0e1117; }
    .brand-header {
        background: linear-gradient(135deg, #1d2a4d 0%, #111930 100%);
        border-radius: 16px; padding: 2rem; margin-bottom: 2rem;
        border-left: 6px solid #f7941d;
        box-shadow: 0 4px 20px rgba(0,0,0,0.2);
    }
    .stDataFrame { border: 1px solid rgba(255,255,255,0.05) !important; border-radius: 12px !important; }
    div.stButton > button:first-child {
        background: linear-gradient(135deg, #f7941d 0%, #e58512 100%) !important;
        color: white !important; border: none !important;
        font-size: 1.1rem !important; font-weight: 700 !important;
        padding: 0.75rem 2rem !important; border-radius: 10px !important;
        box-shadow: 0 4px 15px rgba(247, 148, 29, 0.4) !important;
        transition: all 0.3s ease !important;
    }
    div.stButton > button:first-child:hover {
        transform: translateY(-2px) !important;
        box-shadow: 0 6px 20px rgba(247, 148, 29, 0.6) !important;
    }
    div.stButton > button:first-child:disabled {
        background: rgba(255,255,255,0.1) !important;
        color: rgba(255,255,255,0.3) !important;
        box-shadow: none !important; cursor: not-allowed !important;
    }
    .stTabs [data-baseweb="tab"] { font-weight: 500; padding: 12px 24px; }
</style>
""", unsafe_allow_html=True)

# ==================== SESSION STATE ====================
if "resultados" not in st.session_state:
    st.session_state.resultados = None
if "tempo_execucao" not in st.session_state:
    st.session_state.tempo_execucao = 0

# ==================== CONSTANTES ====================
BRASILAPI_BASE = "https://brasilapi.com.br/api/cnpj/v1"
RECEITAWS_BASE = "https://receitaws.com.br/v1/cnpj"   # fallback

MAX_TENTATIVAS = 3
BACKOFF_BASE = 1.5
TIMEOUT = 25.0
CONCORRENCIA_MAX = 3    # BrasilAPI bloqueia abuso — mantém baixo

STATUS_OK = "OK"
STATUS_NAO_ENCONTRADO = "Não encontrado"
STATUS_ERRO = "Erro"
STATUS_ERRO_FALLBACK = "Erro (fallback usado)"
STATUS_CNPJ_INVALIDO = "CNPJ inválido"

TRANSIENT_STATUS = {408, 425, 429, 500, 502, 503, 504}


# ==================== HELPERS ====================
def limpar_cnpj(texto: str) -> str:
    return re.sub(r"\D", "", texto)


def formatar_cnpj(cnpj: str) -> str:
    c = limpar_cnpj(cnpj)
    if len(c) == 14:
        return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}"
    return cnpj


def formatar_data_iso(iso_str: str | None) -> str:
    """Converte '2024-01-06' → '06/01/2024'. Retorna '---' se vazio."""
    if not iso_str:
        return "---"
    try:
        return date.fromisoformat(iso_str[:10]).strftime("%d/%m/%Y")
    except (ValueError, TypeError):
        return iso_str  # devolve como veio se não puder parsear


def bool_para_texto(valor: Any) -> str:
    """
    Converte o campo opcao_pelo_simples/mei da API para texto exibível.
    - True  → 'Sim'
    - False → 'Não'
    - None  → 'N/D' (não declarado pela Receita)
    """
    if valor is True:
        return "Sim"
    if valor is False:
        return "Não"
    return "N/D"


# ==================== CLIENTE HTTP ====================
async def _get_json(client: httpx.AsyncClient, url: str) -> tuple[dict | None, str | None]:
    """
    GET JSON com retry para erros transitórios.
    Retorna (dados, erro). Se sucesso, erro é None.
    Trata 404 como "not_found" (marcador especial).
    """
    ultimo_erro = None
    for tentativa in range(MAX_TENTATIVAS):
        try:
            resp = await client.get(url)
            if resp.status_code == 404:
                return None, "not_found"
            if resp.status_code == 200:
                try:
                    return resp.json(), None
                except Exception as e:
                    return None, f"JSON inválido: {e}"

            # Erro HTTP
            msg = f"HTTP {resp.status_code}"
            try:
                body = resp.json()
                if isinstance(body, dict):
                    if "message" in body:
                        msg += f": {body['message'][:120]}"
                    elif "detail" in body:
                        msg += f": {body['detail'][:120]}"
            except Exception:
                pass

            if resp.status_code in TRANSIENT_STATUS:
                ultimo_erro = msg
                if tentativa < MAX_TENTATIVAS - 1:
                    await asyncio.sleep(BACKOFF_BASE * (tentativa + 1))
                    continue
            return None, msg

        except httpx.TimeoutException:
            ultimo_erro = "Timeout"
            if tentativa < MAX_TENTATIVAS - 1:
                await asyncio.sleep(BACKOFF_BASE * (tentativa + 1))
        except Exception as e:
            ultimo_erro = f"{type(e).__name__}: {str(e)[:120]}"
            if tentativa < MAX_TENTATIVAS - 1:
                await asyncio.sleep(BACKOFF_BASE * (tentativa + 1))

    return None, ultimo_erro or "Erro desconhecido"


async def consultar_brasilapi(client: httpx.AsyncClient, cnpj: str) -> tuple[dict | None, str | None]:
    return await _get_json(client, f"{BRASILAPI_BASE}/{cnpj}")


async def consultar_receitaws(client: httpx.AsyncClient, cnpj: str) -> tuple[dict | None, str | None]:
    """Fallback com estrutura diferente — usa campo 'simples.optante'."""
    return await _get_json(client, f"{RECEITAWS_BASE}/{cnpj}")


# ==================== NORMALIZAÇÃO ====================
def normalizar_brasilapi(data: dict) -> dict:
    """
    Extrai os campos padronizados do JSON da BrasilAPI (fonte MinhaReceita).
    Schema oficial:
        razao_social, descricao_situacao_cadastral,
        opcao_pelo_simples, data_opcao_pelo_simples, data_exclusao_do_simples,
        opcao_pelo_mei, data_opcao_pelo_mei, data_exclusao_do_mei
    """
    return {
        "Razão Social":     data.get("razao_social", "").strip() or "---",
        "Situação":         (data.get("descricao_situacao_cadastral") or "---").upper(),
        "Simples Nacional": bool_para_texto(data.get("opcao_pelo_simples")),
        "MEI":              bool_para_texto(data.get("opcao_pelo_mei")),
        "Data Opção":       formatar_data_iso(data.get("data_opcao_pelo_simples")),
        "Data Opção MEI":   formatar_data_iso(data.get("data_opcao_pelo_mei")),
        "simples_bool":     data.get("opcao_pelo_simples") is True,
        "mei_bool":         data.get("opcao_pelo_mei") is True,
        "fonte":            "BrasilAPI",
    }


def normalizar_receitaws(data: dict) -> dict:
    """
    Extrai campos do ReceitaWS (schema diferente):
        nome, situacao, simples: {optante, data_opcao, ...}, simei: {optante, ...}
    """
    simples = data.get("simples") if isinstance(data.get("simples"), dict) else {}
    simei = data.get("simei") if isinstance(data.get("simei"), dict) else {}

    def _fmt_data_br(s):
        if not s:
            return "---"
        # ReceitaWS pode retornar "DD/MM/YYYY" ou "YYYY-MM-DD"
        if "/" in s:
            return s
        return formatar_data_iso(s)

    return {
        "Razão Social":     (data.get("nome") or "").strip() or "---",
        "Situação":         (data.get("situacao") or "---").upper(),
        "Simples Nacional": bool_para_texto(simples.get("optante")),
        "MEI":              bool_para_texto(simei.get("optante")),
        "Data Opção":       _fmt_data_br(simples.get("data_opcao")),
        "Data Opção MEI":   _fmt_data_br(simei.get("data_opcao")),
        "simples_bool":     simples.get("optante") is True,
        "mei_bool":         simei.get("optante") is True,
        "fonte":            "ReceitaWS",
    }


# ==================== CONSULTA UNITÁRIA ====================
async def consultar_um(client: httpx.AsyncClient, cnpj_raw: str) -> dict:
    cnpj = limpar_cnpj(cnpj_raw)
    resultado = {
        "CNPJ":             formatar_cnpj(cnpj),
        "Razão Social":     "---",
        "Situação":         "---",
        "Simples Nacional": "---",
        "MEI":              "---",
        "Data Opção":       "---",
        "Data Opção MEI":   "---",
        "Fonte":            "---",
        "Status":           STATUS_OK,
        "Detalhes":         "",
        "simples_bool":     False,
        "mei_bool":         False,
        "cnpj_raw":         cnpj,
    }

    if len(cnpj) != 14:
        resultado["Razão Social"] = "CNPJ Inválido"
        resultado["Status"] = STATUS_CNPJ_INVALIDO
        resultado["Detalhes"] = "CNPJ precisa ter 14 dígitos"
        return resultado

    # ── Tenta BrasilAPI primeiro ──
    data, erro = await consultar_brasilapi(client, cnpj)

    if data:
        norm = normalizar_brasilapi(data)
        resultado.update({
            "Razão Social":     norm["Razão Social"],
            "Situação":         norm["Situação"],
            "Simples Nacional": norm["Simples Nacional"],
            "MEI":              norm["MEI"],
            "Data Opção":       norm["Data Opção"],
            "Data Opção MEI":   norm["Data Opção MEI"],
            "Fonte":            norm["fonte"],
            "simples_bool":     norm["simples_bool"],
            "mei_bool":         norm["mei_bool"],
        })
        return resultado

    if erro == "not_found":
        resultado["Razão Social"] = "Não localizado"
        resultado["Status"] = STATUS_NAO_ENCONTRADO
        return resultado

    # ── Fallback: ReceitaWS ──
    data_fb, erro_fb = await consultar_receitaws(client, cnpj)

    if data_fb and data_fb.get("status") != "ERROR":
        norm = normalizar_receitaws(data_fb)
        resultado.update({
            "Razão Social":     norm["Razão Social"],
            "Situação":         norm["Situação"],
            "Simples Nacional": norm["Simples Nacional"],
            "MEI":              norm["MEI"],
            "Data Opção":       norm["Data Opção"],
            "Data Opção MEI":   norm["Data Opção MEI"],
            "Fonte":            norm["fonte"],
            "simples_bool":     norm["simples_bool"],
            "mei_bool":         norm["mei_bool"],
            "Status":           STATUS_ERRO_FALLBACK,
            "Detalhes":         f"BrasilAPI: {erro} → usou ReceitaWS",
        })
        return resultado

    if erro_fb == "not_found":
        resultado["Razão Social"] = "Não localizado"
        resultado["Status"] = STATUS_NAO_ENCONTRADO
        return resultado

    # Ambos falharam
    resultado["Razão Social"] = "⚠️ Erro na consulta"
    resultado["Simples Nacional"] = "⚠️ Erro"
    resultado["MEI"] = "⚠️ Erro"
    resultado["Status"] = STATUS_ERRO
    resultado["Detalhes"] = f"BrasilAPI: {erro} | ReceitaWS: {erro_fb}"
    return resultado


async def consultar_lote(cnpjs: list[str], progress_callback) -> list[dict]:
    """Consulta em lote com concorrência limitada."""
    sem = asyncio.Semaphore(CONCORRENCIA_MAX)
    resultados: list[dict | None] = [None] * len(cnpjs)
    concluidos = 0

    async with httpx.AsyncClient(
        timeout=TIMEOUT,
        headers={"User-Agent": "VIXPAR-Portal-Fiscal/1.0"},
        follow_redirects=True,
    ) as client:

        async def worker(idx: int, cnpj: str):
            nonlocal concluidos
            async with sem:
                resultado = await consultar_um(client, cnpj)
                resultados[idx] = resultado
                concluidos += 1
                progress_callback(concluidos, len(cnpjs))

        await asyncio.gather(*[worker(i, c) for i, c in enumerate(cnpjs)])

    return [r for r in resultados if r is not None]


def estilizar_linha(row):
    status = row.get("Status", "")
    if status == STATUS_OK:
        if row.get("Simples Nacional") == "Sim" or row.get("MEI") == "Sim":
            return ["background-color: rgba(35, 134, 54, 0.15)"] * len(row)
        return [""] * len(row)
    if status in (STATUS_ERRO, STATUS_ERRO_FALLBACK):
        return ["background-color: rgba(247, 148, 29, 0.18)"] * len(row)
    if status in (STATUS_CNPJ_INVALIDO, STATUS_NAO_ENCONTRADO):
        return ["background-color: rgba(200, 60, 60, 0.12)"] * len(row)
    return [""] * len(row)


# ==================== UI ====================
def _renderizar_ui():
    with st.container():
        col_logo, col_title = st.columns([1, 4])
        with col_logo:
            try:
                st.image("logo_vixpar.png", use_container_width=True)
            except Exception:
                st.markdown("### VIXPAR")
        with col_title:
            st.markdown("""
            <div class="brand-header">
                <h2 style='margin:0; padding:0; color:white;'>Monitoramento Fiscal Avançado</h2>
                <p style='margin:5px 0 0 0; color:#b0bccc; font-size:0.95rem;'>Consulta em lote automatizada — Simples Nacional & Receita Federal</p>
            </div>
            """, unsafe_allow_html=True)

    tab1, tab2 = st.tabs([
        ":material/query_stats: Painel de Consulta",
        ":material/history: Histórico de Consultas"
    ])

    with tab1:
        col_inp, col_side = st.columns([2, 1], gap="large")

        with col_inp:
            st.markdown("##### :material/input: Dados de Entrada")
            texto_cnpjs = st.text_area(
                "Cole uma lista de CNPJs (um por linha)",
                height=150,
                placeholder="00.000.000/0001-00\n11111111000111",
                label_visibility="collapsed"
            )
            uploaded_file = st.file_uploader(
                "Ou faça upload de um arquivo contendo os CNPJs (.txt, .csv)",
                type=["txt", "csv"]
            )

        with col_side:
            st.markdown("##### :material/tune: Instruções")
            st.caption("1. Adicione os CNPJs ao lado colando ou anexando.")
            st.caption("2. Clique no botão destacado abaixo para processar.")
            st.caption("3. Fonte principal: BrasilAPI (MinhaReceita). Fallback: ReceitaWS.")
            st.caption(f"4. Concorrência: {CONCORRENCIA_MAX} req simultâneas (evita bloqueio).")
            if st.button(":material/refresh: Resetar Painel", use_container_width=True):
                st.session_state.resultados = None
                st.rerun()

        cnpjs_raw = []
        if uploaded_file:
            cnpjs_raw.extend([
                l.strip() for l in uploaded_file.read().decode("utf-8", errors="ignore").splitlines()
                if l.strip()
            ])
        if texto_cnpjs.strip():
            cnpjs_raw.extend([l.strip() for l in texto_cnpjs.splitlines() if l.strip()])

        cnpjs_raw = list(dict.fromkeys(cnpjs_raw))

        st.markdown("---")

        botao_desabilitado = len(cnpjs_raw) == 0
        texto_botao = (
            f":material/bolt: PROCESSAR {len(cnpjs_raw)} CNPJs AGORA"
            if cnpjs_raw
            else ":material/lock: INSIRA CNPJS PARA PROCESSAR"
        )

        if st.button(texto_botao, type="primary", use_container_width=True, disabled=botao_desabilitado):
            inicio = time.time()
            bar_text = st.empty()
            prog_bar = st.progress(0)

            def update_progress(atual, total):
                prog_bar.progress(atual / total)
                bar_text.caption(f"Processando: {atual} de {total} CNPJs analisados...")

            resultados = asyncio.run(consultar_lote(cnpjs_raw, update_progress))
            st.session_state.resultados = resultados
            st.session_state.tempo_execucao = time.time() - inicio

            bar_text.empty()
            prog_bar.empty()

        if st.session_state.resultados:
            df = pd.DataFrame(st.session_state.resultados)

            st.markdown("### :material/analytics: Indicadores de Resumo")
            m1, m2, m3, m4, m5 = st.columns(5)
            qtd_erros = int((df["Status"].isin([STATUS_ERRO])).sum())

            with m1:
                st.metric("Total Processado", len(df))
            with m2:
                st.metric("Optantes Simples", int(df["simples_bool"].sum()))
            with m3:
                st.metric("Microempreendedores (MEI)", int(df["mei_bool"].sum()))
            with m4:
                st.metric(
                    "Erros / Reprocessar",
                    qtd_erros,
                    delta=None if qtd_erros == 0 else "atenção",
                    delta_color="inverse",
                )
            with m5:
                st.metric("Tempo Total", f"{st.session_state.tempo_execucao:.2f}s")

            if qtd_erros > 0:
                st.warning(
                    f"⚠️ {qtd_erros} CNPJ(s) tiveram falha total. "
                    "Colunas mostram ⚠️ Erro em vez de assumir 'Não'."
                )
                if st.button(
                    f":material/replay: Reprocessar {qtd_erros} CNPJ(s) com erro",
                    use_container_width=True,
                ):
                    inicio = time.time()
                    indices_erro = df.index[df["Status"] == STATUS_ERRO].tolist()
                    cnpjs_erro = df.loc[indices_erro, "cnpj_raw"].tolist()

                    bar_text = st.empty()
                    prog_bar = st.progress(0)

                    def update_progress(atual, total):
                        prog_bar.progress(atual / total)
                        bar_text.caption(f"Reprocessando: {atual} de {total}...")

                    novos = asyncio.run(consultar_lote(cnpjs_erro, update_progress))

                    atualizados = list(st.session_state.resultados)
                    for idx, novo in zip(indices_erro, novos):
                        atualizados[idx] = novo
                    st.session_state.resultados = atualizados
                    st.session_state.tempo_execucao += time.time() - inicio

                    bar_text.empty()
                    prog_bar.empty()
                    st.rerun()

            st.markdown("---")
            st.markdown("### :material/table_chart: Dados Consolidados")

            col_f1, col_f2 = st.columns([3, 1])
            with col_f1:
                busca = st.text_input(
                    ":material/search: Filtrar resultados na tela:",
                    placeholder="Digite uma Razão Social ou CNPJ...",
                )
            with col_f2:
                mostrar_apenas = st.selectbox(
                    "Exibir:",
                    ["Todos", "Somente OK", "Somente com erro", "Somente optantes"],
                    index=0,
                )

            df_view = df.copy()
            if busca:
                df_view = df_view[
                    df_view["Razão Social"].str.contains(busca, case=False, na=False)
                    | df_view["CNPJ"].str.contains(busca, na=False)
                ]

            if mostrar_apenas == "Somente OK":
                df_view = df_view[df_view["Status"] == STATUS_OK]
            elif mostrar_apenas == "Somente com erro":
                df_view = df_view[df_view["Status"].isin([STATUS_ERRO, STATUS_ERRO_FALLBACK])]
            elif mostrar_apenas == "Somente optantes":
                df_view = df_view[(df_view["simples_bool"]) | (df_view["mei_bool"])]

            colunas_visiveis = ["CNPJ", "Razão Social", "Situação",
                                "Simples Nacional", "MEI",
                                "Data Opção", "Data Opção MEI",
                                "Fonte", "Status", "Detalhes"]
            df_display = df_view[colunas_visiveis]

            st.dataframe(
                df_display.style.apply(estilizar_linha, axis=1),
                use_container_width=True,
                hide_index=True,
            )

            st.markdown("#### :material/download: Exportar Relatórios")
            c_csv, c_xlsx = st.columns(2)
            df_export = df_view[colunas_visiveis]

            with c_csv:
                csv_data = df_export.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")
                st.download_button(
                    label="Salvar Planilha em CSV",
                    data=csv_data,
                    file_name=f"vixpar_relatorio_fiscal_{datetime.now():%Y%m%d_%H%M}.csv",
                    mime="text/csv",
                    use_container_width=True,
                )

            with c_xlsx:
                output_buffer = io.BytesIO()
                with pd.ExcelWriter(output_buffer, engine="openpyxl") as writer:
                    df_export.to_excel(writer, index=False, sheet_name="Fiscal")

                    from openpyxl.styles import PatternFill, Font, Alignment
                    from openpyxl.utils import get_column_letter
                    ws = writer.sheets["Fiscal"]

                    header_fill = PatternFill("solid", fgColor="1D2A4D")
                    for cell in ws[1]:
                        cell.fill = header_fill
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.alignment = Alignment(horizontal="center", vertical="center")

                    verde = PatternFill("solid", fgColor="D4EDDA")
                    ambar = PatternFill("solid", fgColor="FFE8CC")
                    vermelho = PatternFill("solid", fgColor="F8D7DA")

                    idx_status = colunas_visiveis.index("Status")
                    idx_sn = colunas_visiveis.index("Simples Nacional")
                    idx_mei = colunas_visiveis.index("MEI")

                    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                        s = row[idx_status].value
                        if s == STATUS_OK:
                            if row[idx_sn].value == "Sim" or row[idx_mei].value == "Sim":
                                for c in row:
                                    c.fill = verde
                        elif s in (STATUS_ERRO, STATUS_ERRO_FALLBACK):
                            for c in row:
                                c.fill = ambar
                        elif s in (STATUS_CNPJ_INVALIDO, STATUS_NAO_ENCONTRADO):
                            for c in row:
                                c.fill = vermelho

                    for col in ws.columns:
                        max_len = max(len(str(c.value or "")) for c in col)
                        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 60)

                st.download_button(
                    label="Salvar Planilha em Excel (XLSX)",
                    data=output_buffer.getvalue(),
                    file_name=f"vixpar_relatorio_fiscal_{datetime.now():%Y%m%d_%H%M}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )

    with tab2:
        st.markdown("### :material/history_toggle_off: Histórico Recente")
        if st.session_state.resultados:
            st.caption("Abaixo constam os dados da última execução armazenados em cache temporário de sessão.")
            df_hist = pd.DataFrame(st.session_state.resultados)
            colunas_hist = ["CNPJ", "Razão Social", "Situação",
                            "Simples Nacional", "MEI",
                            "Data Opção", "Data Opção MEI",
                            "Fonte", "Status", "Detalhes"]
            st.dataframe(
                df_hist[colunas_hist].style.apply(estilizar_linha, axis=1),
                use_container_width=True,
                hide_index=True,
            )
        else:
            st.info("Nenhuma consulta em lote executada nesta sessão.")

    st.markdown("---")
    st.markdown(
        f"<p style='text-align: center; color: #55637a; font-size: 0.85rem;'>© {datetime.now().year} VIXPAR — Setor de Inteligência e Compliance Fiscal</p>",
        unsafe_allow_html=True
    )


# ==================== ENTRY POINT ====================
try:
    from streamlit.runtime import exists as _st_runtime_exists
    if _st_runtime_exists():
        _renderizar_ui()
except Exception:
    pass
